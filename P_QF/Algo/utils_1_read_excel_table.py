# -*- coding: utf-8 -*-
"""
@author: lyu
"""
import pandas as pd
import openpyxl

# Extract Excel Table to DataFrame
def read_table(file_name: str, table_name: str) -> pd.DataFrame:
    wb = openpyxl.load_workbook(file_name, read_only=False, data_only=True)
    for sheetname in wb.sheetnames:
        sheet = wb[sheetname]
        if table_name in sheet.tables:
            tbl = sheet.tables[table_name]
            tbl_range = tbl.ref
            break
    data = sheet[tbl_range]
    content = [[cell.value for cell in row] for row in data]
    header = content[0]
    rest = content[1:]
    df = pd.DataFrame(rest, columns=header)
    wb.close()
    return df

if __name__ == "__main__":
    Logit = read_table(r'C:\Users\lyu\Desktop\Spreadsheets\Quant_Logit_and_Probit.xlsx', "Table1")
    print(Logit)